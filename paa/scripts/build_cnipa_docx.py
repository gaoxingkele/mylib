# -*- coding: utf-8 -*-
"""把一案的 02/03/04/05 Markdown 装配成 CNIPA 版式的完整申请文件 Word。

- 五部分各自起页：权利要求书、说明书、说明书附图、说明书摘要、摘要附图
- 公式：LaTeX（\\( \\) / \\[ \\] / $ $）经 pandoc 转为 Word 原生公式（OMML，可编辑）
- 附图：由 04 的 mermaid 源码按 neutral 主题重绘为黑白线条 PNG（需 npx @mermaid-js/mermaid-cli）；
  重绘失败时回退 `<ver_dir>/附图/图_NNN.png` 或 `fig_NNN.png`
- 版式：A4；页边距 上25 下15 左25 右15 mm；宋体小四、1.5 倍行距、首行缩进 2 字符；页脚页码；标题黑体

输入目录约定（<ver_dir>）：
  02_权利要求书.md   抬头 `## 发明名称：...`，权项写作 `**1.** ...`
  03_说明书.md       二级标题为章节；`# PAA 附注` 之后的内容不进 Word
  04_附图清单与描述.md  每图 `### 图N：标题` + ```mermaid 块
  05_说明书摘要.md   含 `摘要附图：图N` 行

用法：
  python build_cnipa_docx.py <ver_dir> [--out <docx>] [--case-name <名>] [--fig-dir <黑白附图输出目录>]
依赖：pypandoc（含 pandoc）、python-docx、Pillow、mermaid-cli。
来源：zhuanlishenqing 电力三案 scripts/build_dcr3_word.py 的通用化（2026-10-13）。
"""
import argparse, json, os, re, shutil, subprocess, sys, tempfile

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

PAGEBREAK = '\n\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\n'
MERMAID_CFG = {"theme": "neutral", "themeVariables": {
    "background": "#ffffff", "primaryColor": "#ffffff", "primaryBorderColor": "#000000",
    "lineColor": "#000000", "primaryTextColor": "#000000", "secondaryColor": "#ffffff",
    "tertiaryColor": "#ffffff", "clusterBkg": "#ffffff", "clusterBorder": "#000000"}}


def read(p):
    return open(p, encoding="utf-8").read()


def strip_head(md, drop_title=True):
    """去掉抬头（# 标题、## 案件、## 发明名称、版本说明、分隔线），返回(发明名称, 正文)。"""
    name, out = "", []
    for line in md.splitlines():
        s = line.strip()
        m = re.match(r"^#+\s*发明名称[：:]\s*(.+)$", s)
        if m:
            name = m.group(1).strip(); continue
        if drop_title and re.match(r"^#\s*(权利要求书|说明书|说明书摘要)\s*$", s):
            continue
        if re.match(r"^#+\s*案件[：:]", s) or re.match(r"^（版本\s*[\d.]+）$", s) or s == "---":
            continue
        out.append(line)
    return name, "\n".join(out).strip()


def render_figures(ver_dir, work):
    md = read(os.path.join(ver_dir, "04_附图清单与描述.md"))
    figs = []
    for m in re.finditer(r"^###\s*图\s*(\d+)[：:]\s*(.+?)\s*$\n(?:.*\n)*?```mermaid\n(.*?)```", md, re.M | re.S):
        n, title, src = int(m.group(1)), m.group(2).strip(), m.group(3)
        mmd, png = os.path.join(work, f"fig{n:02d}.mmd"), os.path.join(work, f"fig{n:02d}.png")
        open(mmd, "w", encoding="utf-8").write(src)
        cfg = os.path.join(work, "mermaid.json"); json.dump(MERMAID_CFG, open(cfg, "w"))
        r = subprocess.run(f'npx -y @mermaid-js/mermaid-cli -i "{mmd}" -o "{png}" -b white -s 3 -c "{cfg}"',
                           shell=True, capture_output=True, text=True)
        if r.returncode != 0 or not os.path.exists(png):
            for cand in (os.path.join(ver_dir, "附图", f"图_{n:03d}.png"), os.path.join(ver_dir, "附图", f"fig_{n:03d}.png")):
                if os.path.exists(cand):
                    shutil.copy(cand, png); break
            print(f"  [warn] 图{n} mermaid 重绘失败，回退已有 PNG: {os.path.exists(png)}; {r.stderr[-200:]}")
        figs.append((n, title, png))
    return figs


def img_size(p, max_w=15.0, max_h=20.0):
    from PIL import Image
    try:
        w, h = Image.open(p).size
    except Exception:
        return "width=15cm"
    return f"width={max_w:.1f}cm" if h / w * max_w <= max_h else f"height={max_h:.1f}cm"


def assemble(ver_dir):
    work = tempfile.mkdtemp(prefix="cnipa_docx_")
    name, claims = strip_head(read(os.path.join(ver_dir, "02_权利要求书.md")))
    _, spec = strip_head(read(os.path.join(ver_dir, "03_说明书.md")))
    _, abst = strip_head(read(os.path.join(ver_dir, "05_说明书摘要.md")))
    m = re.search(r"摘要附图[：:]\s*图\s*(\d+)", abst); absfig = int(m.group(1)) if m else 1
    abst = re.sub(r"^\s*摘要附图[：:].*$", "", abst, flags=re.M).strip()
    abst = re.split(r"\n#+\s", "\n" + abst)[0].strip() or abst
    figs = render_figures(ver_dir, work)
    spec = re.sub(r"^## ", "### ", spec, flags=re.M)
    claims = re.sub(r"\*\*(\d+)\.\*\*", r"\1\\.", claims)
    claims = re.sub(r"^(\d+)\.\s", r"\1\\. ", claims, flags=re.M)
    spec = re.split(r"^#+\s*PAA\s*附注.*$", spec, flags=re.M)[0].rstrip()
    spec = re.sub(r"（v\d+\.\d+\s*增补）|（版本\s*\d+\.\d+）", "", spec)
    fix_tag = lambda s: re.sub(r"\\tag\*?\{([^}]*)\}", r"\\qquad(\1)", s)
    claims, spec = fix_tag(claims), fix_tag(spec)
    parts = [f"# 权利要求书\n\n{claims}",
             f"# 说明书\n\n::: {{custom-style=\"InventionTitle\"}}\n{name}\n:::\n\n{spec}",
             "# 说明书附图\n\n" + PAGEBREAK.join(
                 f"![]({p.replace(os.sep, '/')}){{{img_size(p)}}}\n\n::: {{custom-style=\"FigCaption\"}}\n图{n}\n:::" for n, t, p in figs),
             f"# 说明书摘要\n\n{abst}"]
    af = [p for n, t, p in figs if n == absfig]
    parts.append("# 摘要附图\n\n" + (f"![]({af[0].replace(os.sep, '/')}){{{img_size(af[0])}}}" if af else ""))
    return name, PAGEBREAK.join(parts), work, figs, absfig


def set_run_font(run, east="宋体", west="Times New Roman", size=12, bold=None):
    run.font.name = west; run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), east); rf.set(qn("w:ascii"), west); rf.set(qn("w:hAnsi"), west)


def add_page_number(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run()
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._element.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._element.append(it)
        set_run_font(r, size=10.5)


def format_docx(path):
    d = Document(path)
    for sn in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Title"):
        try:
            d.styles[sn].font.color.rgb = None
            rpr = d.styles[sn].element.get_or_add_rPr(); c = rpr.find(qn("w:color"))
            if c is not None: rpr.remove(c)
        except KeyError:
            pass
    for s in d.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Mm(25), Mm(15), Mm(25), Mm(15)
        add_page_number(s)
    for p in d.paragraphs:
        st = p.style.name if p.style is not None else ""
        has_math = p._element.find(".//" + qn("m:oMathPara")) is not None
        has_img = p._element.find(".//" + qn("w:drawing")) is not None
        pf = p.paragraph_format
        if st.startswith("Heading 1"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: set_run_font(r, east="黑体", size=16, bold=True)
            pf.space_after = Pt(12); continue
        if st.startswith("Heading"):
            for r in p.runs: set_run_font(r, east="黑体", size=12, bold=True)
            pf.space_before = Pt(6); pf.space_after = Pt(6); pf.line_spacing = 1.5; continue
        if st == "InventionTitle":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: set_run_font(r, east="黑体", size=14, bold=True)
            pf.space_after = Pt(12); continue
        if st == "FigCaption" or has_img:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: set_run_font(r, size=12)
            continue
        if has_math:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER; continue
        pf.line_spacing = 1.5; pf.first_line_indent = Pt(24); pf.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs: set_run_font(r, size=12)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs: set_run_font(r, size=10.5)
    d.core_properties.author = ""; d.core_properties.last_modified_by = ""; d.core_properties.comments = ""
    d.save(path)


def build(ver_dir, out=None, case_name=None, fig_dir=None):
    ver_dir = os.path.abspath(ver_dir)
    name, md, work, figs, absfig = assemble(ver_dir)
    tag = os.path.basename(ver_dir).split("_")[-1]
    case_name = case_name or os.path.basename(os.path.dirname(ver_dir))
    out = out or os.path.join(ver_dir, f"{case_name}_完整申请文件_{tag}_CNIPA版式.docx")
    pypandoc.convert_text(md, "docx", format="markdown+tex_math_single_backslash+tex_math_dollars+raw_attribute+fenced_divs+pipe_tables",
                          outputfile=out, extra_args=[f"--resource-path={work}"])
    format_docx(out)
    fig_dir = fig_dir or os.path.join(ver_dir, "附图_黑白")
    os.makedirs(fig_dir, exist_ok=True)
    for n, t, p in figs:
        if os.path.exists(p):
            shutil.copy(p, os.path.join(fig_dir, f"图{n}.png"))
    shutil.rmtree(work, ignore_errors=True)
    print(f"[ok] {out}  附图 {len(figs)} 幅，摘要附图 图{absfig}")
    return out


def check(path):
    """粗检：无残留 LaTeX 命令、五部分标题齐全、图片数。"""
    import zipfile
    z = zipfile.ZipFile(path); x = z.read("word/document.xml").decode("utf8")
    txt = re.sub(r"<[^>]+>", "", x)
    imgs = len([n for n in z.namelist() if n.startswith("word/media/")])
    latex = len(re.findall(r"\\(frac|sum|mathrm|left|right|cdot|le|ge|times|text|alpha|beta|hat)(?![a-z])", txt)) + txt.count("$$")
    heads = [h for h in ("权利要求书", "说明书附图", "说明书摘要", "摘要附图", "技术领域", "背景技术", "具体实施方式") if h in txt]
    print(f"[check] chars={len(txt)} images={imgs} latex_residue={latex} sections={len(heads)}/7")
    return latex == 0 and len(heads) == 7


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("ver_dir", help="含 02/03/04/05 的版本目录，如 <案>/申请书_v2.8")
    ap.add_argument("--out"); ap.add_argument("--case-name"); ap.add_argument("--fig-dir")
    ap.add_argument("--check", action="store_true", help="生成后做粗检")
    a = ap.parse_args()
    o = build(a.ver_dir, a.out, a.case_name, a.fig_dir)
    if a.check:
        sys.exit(0 if check(o) else 1)
